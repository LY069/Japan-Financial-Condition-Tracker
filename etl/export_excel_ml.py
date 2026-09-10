"""Build the refreshable Monetary & Liquidity workbook: Japan_ML_Tracker.xlsx.

Sheets: Dashboard (MCI/LCI/MLCI + 7 axes + assessment), Monetary, Liquidity,
Data (wide), Indicators, Catalog. Regenerated from the SQLite DB each run.
"""
from __future__ import annotations

import os
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from db import connect, get_series, catalog_rows, get_meta
from build_indicators import month_grid, to_monthly
from export_excel import (fill, score_fill, style_header_row, title_banner, BORDER, H2, BOLD,
                          NAVY, GREY, LIGHT)
from export_web import label_for, latest_ind, latest_obs
from export_web_ml import assessment
from ml_framework import ML_AXES, ML_AXIS_ORDER, ML_STAGES

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "Japan_ML_Tracker.xlsx")


def members():
    return sorted({sid for a in ML_AXES.values() for sid in a["members"]})


def build_dashboard(wb, conn):
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False
    title_banner(ws, "Japan Monetary & Liquidity Conditions Tracker", 6)
    mlci = latest_ind(conn, "composite", "mlci")
    ws["A3"] = "As of"; ws["A3"].font = BOLD; ws["B3"] = mlci["date"] if mlci else ""
    ws["D3"] = "Data mode"; ws["D3"].font = BOLD; ws["E3"] = get_meta(conn, "data_mode", "")

    def lv(sid):
        _, v = latest_obs(conn, sid)
        return v

    cards = []
    for ck, lab in [("mci", "Monetary conditions (MCI)"), ("lci", "Liquidity conditions (LCI)"),
                    ("mlci", "Combined (MLCI)")]:
        li = latest_ind(conn, "composite", ck)
        cards.append((lab, f"{li['score']:+.2f}  {label_for(li['score'])}" if li else "n/a"))
    cards += [("Policy rate", f"{lv('policy_rate'):.2f}%"), ("M2 y/y", f"{lv('m2_yoy'):+.1f}%"),
              ("REER (2020=100)", f"{lv('reer'):.1f}")]
    r = 5
    for i, (lab, val) in enumerate(cards):
        col, row = 1 + (i % 3) * 2, r + (i // 3) * 3
        ws.cell(row=row, column=col, value=lab).font = Font(bold=True, color=GREY, size=9)
        c = ws.cell(row=row + 1, column=col, value=val); c.font = Font(bold=True, size=13, color=NAVY)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)

    ar = 12
    ws.cell(row=ar, column=1, value="Framework axes").font = H2
    ar += 1
    for j, h in enumerate(["Axis", "Stage", "Score (z)", "Assessment"], start=1):
        ws.cell(row=ar, column=j, value=h)
    style_header_row(ws, ar, 4)
    for a in ML_AXIS_ORDER:
        li = latest_ind(conn, "axis", f"ml::{a}")
        sc = li["score"] if li else None
        ar += 1
        ws.cell(row=ar, column=1, value=ML_AXES[a]["label"]).border = BORDER
        ws.cell(row=ar, column=2, value=ML_STAGES[ML_AXES[a]["stage"]]).border = BORDER
        c = ws.cell(row=ar, column=3, value=round(sc, 2) if sc is not None else None)
        c.border = BORDER; c.fill = score_fill(sc); c.alignment = Alignment(horizontal="center")
        c4 = ws.cell(row=ar, column=4, value=label_for(sc)); c4.border = BORDER; c4.fill = score_fill(sc)

    ar += 2
    ws.cell(row=ar, column=1, value="Assessment").font = H2
    ar += 1
    h = {}
    for ck in ("mci", "lci", "mlci"):
        li = latest_ind(conn, "composite", ck)
        h[f"{ck}_score"] = li["score"] if li else 0.0
    for sid in ["m2_yoy", "boj_assets_yoy", "monetary_base_yoy", "reer", "tibor_ois_3m",
                "jpy_basis_3m", "jgb_market_functioning_di", "boj_jgb_share"]:
        h[sid] = lv(sid) or 0.0
    ws.merge_cells(start_row=ar, start_column=1, end_row=ar + 5, end_column=6)
    tc = ws.cell(row=ar, column=1, value=assessment(h))
    tc.alignment = Alignment(wrap_text=True, vertical="top"); tc.fill = fill(LIGHT)

    # composite history chart (helper data in hidden columns)
    hc = 9
    rows = conn.execute("SELECT date,score FROM indicators WHERE scope='composite' AND key='mci' ORDER BY date").fetchall()
    lci = {r["date"]: r["score"] for r in conn.execute(
        "SELECT date,score FROM indicators WHERE scope='composite' AND key='lci'")}
    ws.cell(row=1, column=hc, value="date"); ws.cell(row=1, column=hc + 1, value="MCI"); ws.cell(row=1, column=hc + 2, value="LCI")
    for i, rr in enumerate(rows, start=2):
        ws.cell(row=i, column=hc, value=rr["date"]); ws.cell(row=i, column=hc + 1, value=round(rr["score"], 3))
        if rr["date"] in lci:
            ws.cell(row=i, column=hc + 2, value=round(lci[rr["date"]], 3))
    ch = LineChart(); ch.title = "MCI vs LCI (z; + = easier / more liquid)"; ch.height = 7; ch.width = 18
    ch.add_data(Reference(ws, min_col=hc + 1, max_col=hc + 2, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=hc, min_row=2, max_row=len(rows) + 1))
    ws.add_chart(ch, f"A{ar + 7}")
    for col, w in {"A": 32, "B": 22, "C": 14, "D": 22, "E": 18, "F": 12}.items():
        ws.column_dimensions[col].width = w
    for c in (hc, hc + 1, hc + 2):
        ws.column_dimensions[get_column_letter(c)].hidden = True


def build_stage_sheet(wb, conn, title, stage):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    title_banner(ws, f"{ML_STAGES[stage]} — indicators", 7)
    heads = ["Axis", "Indicator", "Latest date", "Latest value", "Unit", "Weight", "Score (z)"]
    for j, hd in enumerate(heads, start=1):
        ws.cell(row=3, column=j, value=hd)
    style_header_row(ws, 3, len(heads))
    cat = {r["series_id"]: r for r in catalog_rows(conn)}
    r = 4
    for a in ML_AXIS_ORDER:
        if ML_AXES[a]["stage"] != stage:
            continue
        for sid, w in ML_AXES[a]["members"].items():
            meta = cat[sid]
            d, v = latest_obs(conn, sid)
            li = latest_ind(conn, "series", f"ml::{sid}")
            sc = li["score"] if li else None
            vals = [ML_AXES[a]["label"], meta["name"], d, round(v, 3) if v is not None else None,
                    meta["unit"], w, round(sc, 2) if sc is not None else None]
            for j, val in enumerate(vals, start=1):
                c = ws.cell(row=r, column=j, value=val); c.border = BORDER
                if j == 7:
                    c.fill = score_fill(sc); c.alignment = Alignment(horizontal="center")
            r += 1
    for col, w in zip("ABCDEFG", [26, 44, 12, 13, 12, 8, 10]):
        ws.column_dimensions[col].width = w


def build_data_wide(wb, conn):
    ws = wb.create_sheet("Data (wide)")
    sids = members() + ["policy_rate", "core_cpi_yoy"]
    grid = month_grid(start="1974-01", end=None)
    ws.cell(row=1, column=1, value="date")
    for j, sid in enumerate(sids, start=2):
        ws.cell(row=1, column=j, value=sid)
    style_header_row(ws, 1, len(sids) + 1, color=NAVY)
    monthly = {sid: to_monthly(get_series(conn, sid), grid) for sid in sids}
    for i, d in enumerate(grid, start=2):
        ws.cell(row=i, column=1, value=d)
        for j, sid in enumerate(sids, start=2):
            v = monthly[sid].get(d)
            if v is not None:
                ws.cell(row=i, column=j, value=round(v, 4))
    ws.freeze_panes = "B2"; ws.column_dimensions["A"].width = 12


def build_indicators_sheet(wb, conn):
    ws = wb.create_sheet("Indicators")
    keys = [("composite", "mci", "MCI"), ("composite", "lci", "LCI"), ("composite", "mlci", "MLCI"),
            ("stage", "ml_monetary", "ml_monetary"), ("stage", "ml_liquidity", "ml_liquidity")]
    keys += [("axis", f"ml::{a}", a) for a in ML_AXIS_ORDER]
    grid = month_grid(start="1974-01", end=None)
    ws.cell(row=1, column=1, value="date")
    for j, (_, _, lab) in enumerate(keys, start=2):
        ws.cell(row=1, column=j, value=lab)
    style_header_row(ws, 1, len(keys) + 1, color=NAVY)
    data = {(s, k): {r["date"]: r["score"] for r in conn.execute(
        "SELECT date,score FROM indicators WHERE scope=? AND key=?", (s, k))} for s, k, _ in keys}
    for i, d in enumerate(grid, start=2):
        ws.cell(row=i, column=1, value=d)
        for j, (s, k, _) in enumerate(keys, start=2):
            v = data[(s, k)].get(d)
            if v is not None:
                ws.cell(row=i, column=j, value=round(v, 4))
    ws.freeze_panes = "B2"; ws.column_dimensions["A"].width = 12


def build_catalog_sheet(wb, conn):
    ws = wb.create_sheet("Catalog")
    cols = ["series_id", "name", "stage", "category", "unit", "frequency", "polarity", "source",
            "source_series_id", "source_url", "notes"]
    for j, c in enumerate(cols, start=1):
        ws.cell(row=1, column=j, value=c)
    style_header_row(ws, 1, len(cols), color=NAVY)
    i = 2
    ms = set(members())
    for r in catalog_rows(conn):
        if r["series_id"] in ms:
            for j, c in enumerate(cols, start=1):
                ws.cell(row=i, column=j, value=r[c])
            i += 1
    for col, w in zip("ABCDEFGHIJK", [22, 44, 14, 18, 12, 10, 9, 10, 18, 40, 60]):
        ws.column_dimensions[col].width = w


def main():
    conn = connect()
    wb = Workbook()
    build_dashboard(wb, conn)
    build_stage_sheet(wb, conn, "Monetary", "ml_monetary")
    build_stage_sheet(wb, conn, "Liquidity", "ml_liquidity")
    build_data_wide(wb, conn)
    build_indicators_sheet(wb, conn)
    build_catalog_sheet(wb, conn)
    wb.save(OUT)
    print(f"Wrote {OUT} ({os.path.getsize(OUT)//1024} KB).")


if __name__ == "__main__":
    main()
